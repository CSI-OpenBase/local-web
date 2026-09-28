"""Export the minimal XLSX work archive consumed by CSI Core."""

from __future__ import annotations

from datetime import datetime
from io import BytesIO
from typing import Any, Iterable, Mapping
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill


BEIJING = ZoneInfo("Asia/Shanghai")
ARCHIVE_HEADERS = ("视频 ID", "标题", "发布时间")


def _beijing_time(value: object) -> str:
    raw = str(value or "").strip()
    if not raw:
        return ""
    try:
        parsed = datetime.fromisoformat(raw.replace("Z", "+00:00"))
    except ValueError:
        return raw
    if parsed.tzinfo is None:
        parsed = parsed.replace(tzinfo=BEIJING)
    return parsed.astimezone(BEIJING).strftime("%Y-%m-%d %H:%M:%S")


def build_csi_work_archive_rows(
    videos: Iterable[Mapping[str, Any]],
) -> list[tuple[str, str, str]]:
    """Build unique, privacy-minimal rows in reverse publication order."""

    rows: list[tuple[str, str, str]] = []
    seen_ids: set[str] = set()
    for video in videos:
        video_id = str(video.get("video_id") or "").strip()
        if not video_id or video_id in seen_ids:
            continue
        record = video.get("record")
        source = record if isinstance(record, Mapping) else {}
        platform = str(video.get("platform") or source.get("platform") or "douyin").strip()
        if platform != "douyin":
            continue

        title = str(
            source.get("desc")
            or source.get("description")
            or video.get("title")
            or source.get("title")
            or ""
        ).strip()
        published_at = _beijing_time(source.get("published_at"))
        rows.append((video_id, title, published_at))
        seen_ids.add(video_id)

    rows.sort(key=lambda row: (row[2], row[0]), reverse=True)
    return rows


def build_csi_work_archive(videos: Iterable[Mapping[str, Any]]) -> bytes:
    """Return an XLSX containing exactly 视频 ID, 标题 and 发布时间."""

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "视频档案"
    sheet.freeze_panes = "A2"
    sheet.append(ARCHIVE_HEADERS)

    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="20242A")
        cell.alignment = Alignment(vertical="center")

    for row in build_csi_work_archive_rows(videos):
        sheet.append(row)
        for cell in sheet[sheet.max_row]:
            cell.data_type = "s"

    sheet.column_dimensions["A"].width = 24
    sheet.column_dimensions["B"].width = 64
    sheet.column_dimensions["C"].width = 22
    sheet.auto_filter.ref = f"A1:C{sheet.max_row}"

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
