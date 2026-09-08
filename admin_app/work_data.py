"""Canonical work-performance snapshots and creator-center file ingestion."""

from __future__ import annotations

import csv
import hashlib
import json
import math
import os
import re
import tempfile
import zipfile
from datetime import date, datetime, timezone
from pathlib import Path
from typing import Any, Iterable, Iterator, Mapping


WORK_SCHEMA_VERSION = 1
MAX_XLSX_ARCHIVE_MEMBERS = 1_000
MAX_XLSX_UNCOMPRESSED_BYTES = 128 * 1024 * 1024
MAX_CREATOR_EXPORT_ROWS = 50_000
HEADER_ALIASES = {
    "work_id": "work_id",
    "作品ID": "work_id",
    "视频ID": "work_id",
    "item_id": "work_id",
    "aweme_id": "work_id",
    "video_id": "work_id",
    "itemId": "work_id",
    "awemeId": "work_id",
    "作品编号": "work_id",
    "视频编号": "work_id",
    "作品名称": "title",
    "作品标题": "title",
    "视频标题": "title",
    "标题": "title",
    "title": "title",
    "发布时间": "published_at",
    "publish_time": "published_at",
    "体裁": "content_type",
    "内容类型": "content_type",
    "审核状态": "audit_status",
    "audit_status": "audit_status",
    "播放量": "view_count",
    "播放": "view_count",
    "views": "view_count",
    "完播率": "completion_rate",
    "completion_rate": "completion_rate",
    "wtr": "completion_rate",
    "5s完播率": "five_second_completion_rate",
    "5秒完播率": "five_second_completion_rate",
    "retention_5s": "five_second_completion_rate",
    "wtr_5s": "five_second_completion_rate",
    "封面点击率": "cover_click_rate",
    "cover_ctr": "cover_click_rate",
    "ctr": "cover_click_rate",
    "2s跳出率": "two_second_bounce_rate",
    "2秒跳出率": "two_second_bounce_rate",
    "bounce_2s": "two_second_bounce_rate",
    "平均播放时长": "average_watch_seconds",
    "avg_watch_time": "average_watch_seconds",
    "点赞量": "like_count",
    "点赞": "like_count",
    "likes": "like_count",
    "分享量": "share_count",
    "分享": "share_count",
    "shares": "share_count",
    "评论量": "comment_count",
    "评论": "comment_count",
    "comments": "comment_count",
    "收藏量": "collect_count",
    "收藏": "collect_count",
    "favorites": "collect_count",
    "favs": "collect_count",
    "主页访问量": "profile_visit_count",
    "主页访问": "profile_visit_count",
    "profile_visits": "profile_visit_count",
    "pv": "profile_visit_count",
    "粉丝增量": "follower_gain",
    "涨粉量": "follower_gain",
    "粉丝增长": "follower_gain",
    "followers_gain": "follower_gain",
    "标签": "tags",
    "作品标签": "tags",
    "视频标签": "tags",
    "内容标签": "tags",
    "话题标签": "tags",
    "话题": "tags",
    "tags": "tags",
}
COUNT_FIELDS = (
    "view_count",
    "like_count",
    "share_count",
    "comment_count",
    "collect_count",
    "profile_visit_count",
    "follower_gain",
)
RATE_FIELDS = (
    "completion_rate",
    "five_second_completion_rate",
    "cover_click_rate",
    "two_second_bounce_rate",
)
WORK_ID_RE = re.compile(r"^[A-Za-z0-9][A-Za-z0-9._:-]{0,127}$")


class WorkDataError(ValueError):
    """Raised when a work-performance source violates the canonical contract."""


class _MissingWorkHeader(WorkDataError):
    """Raised when one tabular source has no recognizable work-data header."""


def utc_now() -> str:
    return datetime.now(timezone.utc).isoformat(timespec="microseconds").replace(
        "+00:00", "Z"
    )


def _normalize_header(value: Any) -> str:
    return re.sub(r"\s+", "", str(value or "").lstrip("\ufeff")).casefold()


NORMALIZED_HEADERS = {
    _normalize_header(label): field for label, field in HEADER_ALIASES.items()
}


def _text(value: Any) -> str:
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value)).strip()


def _iso_timestamp(value: Any, *, field: str, allow_empty: bool = False) -> str | None:
    if value in (None, "", "-"):
        if allow_empty:
            return None
        raise WorkDataError(f"{field} is required")
    if isinstance(value, datetime):
        parsed = value
    elif isinstance(value, date):
        parsed = datetime(value.year, value.month, value.day)
    else:
        candidate = _text(value).replace("/", "-")
        try:
            parsed = datetime.fromisoformat(candidate.replace("Z", "+00:00"))
        except ValueError as exc:
            raise WorkDataError(f"{field} is not a valid timestamp: {value!r}") from exc
    if parsed.tzinfo is None:
        parsed = parsed.replace(tzinfo=timezone.utc)
    return parsed.astimezone(timezone.utc).isoformat().replace("+00:00", "Z")


def _number(value: Any, *, field: str, integer: bool, allow_empty: bool) -> int | float | None:
    if value in (None, "", "-"):
        return None if allow_empty else 0
    if isinstance(value, bool):
        raise WorkDataError(f"{field} must be numeric")
    candidate = value
    if isinstance(value, str):
        candidate = value.strip().replace(",", "")
        if candidate.endswith("%"):
            candidate = candidate[:-1]
    try:
        result = float(candidate)
    except (TypeError, ValueError) as exc:
        raise WorkDataError(f"{field} must be numeric") from exc
    if not math.isfinite(result) or result < 0:
        raise WorkDataError(f"{field} must be a non-negative finite number")
    if integer:
        if not result.is_integer():
            raise WorkDataError(f"{field} must be an integer")
        return int(result)
    return result


def _rate(value: Any, *, field: str) -> float | None:
    result = _number(value, field=field, integer=False, allow_empty=True)
    if result is None:
        return None
    rate = float(result)
    if rate > 1:
        rate /= 100
    if rate > 1:
        raise WorkDataError(f"{field} must be between 0 and 1 (or 0 and 100%)")
    return round(rate, 8)


def _tags(value: Any) -> list[str]:
    if value in (None, "", "-"):
        return []
    if isinstance(value, list):
        candidates = value
    else:
        candidates = re.split(r"[,，;；\s]+", str(value))
    return sorted({_text(item).lstrip("#") for item in candidates if _text(item)})


def synthetic_work_id(platform: str, title: str, published_at: str | None) -> str:
    digest = hashlib.sha256(
        "\x1f".join((platform, title.casefold(), published_at or "")).encode("utf-8")
    ).hexdigest()
    return f"work_{digest[:32]}"


def normalize_work_snapshot(
    raw: Mapping[str, Any],
    *,
    platform: str = "douyin",
    observed_at: str | None = None,
    source_file: str = "",
    source_sheet: str = "",
) -> dict[str, Any]:
    title = _text(raw.get("title"))
    if not title:
        raise WorkDataError("title is required")
    normalized_platform = _text(raw.get("platform") or platform).lower()
    if not re.fullmatch(r"[a-z][a-z0-9_-]{1,31}", normalized_platform):
        raise WorkDataError("platform is invalid")
    published_at = _iso_timestamp(
        raw.get("published_at"), field="published_at", allow_empty=True
    )
    observed = _iso_timestamp(
        raw.get("observed_at") or observed_at or utc_now(),
        field="observed_at",
    )
    supplied_id = _text(raw.get("work_id"))
    if supplied_id and not WORK_ID_RE.fullmatch(supplied_id):
        raise WorkDataError("work_id contains unsupported characters")
    work_id = supplied_id or synthetic_work_id(
        normalized_platform, title, published_at
    )
    result: dict[str, Any] = {
        "schema_version": WORK_SCHEMA_VERSION,
        "platform": normalized_platform,
        "work_id": work_id,
        "work_id_kind": "platform" if supplied_id else "synthetic",
        "title": title,
        "tags": _tags(raw.get("tags")),
        "published_at": published_at,
        "content_type": _text(raw.get("content_type")),
        "audit_status": _text(raw.get("audit_status")),
    }
    for field in COUNT_FIELDS:
        result[field] = _number(
            raw.get(field), field=field, integer=True, allow_empty=False
        )
    for field in RATE_FIELDS:
        result[field] = _rate(raw.get(field), field=field)
    average_watch = _number(
        raw.get("average_watch_seconds"),
        field="average_watch_seconds",
        integer=False,
        allow_empty=True,
    )
    result.update(
        {
            "average_watch_seconds": (
                round(float(average_watch), 6) if average_watch is not None else None
            ),
            "observed_at": observed,
            "source_file": Path(source_file).name if source_file else "",
            "source_sheet": _text(source_sheet),
        }
    )
    return result


def _header_map(row: Iterable[Any]) -> dict[int, str]:
    result: dict[int, str] = {}
    used: set[str] = set()
    for index, value in enumerate(row):
        field = NORMALIZED_HEADERS.get(_normalize_header(value))
        if field and field not in used:
            result[index] = field
            used.add(field)
    return result


def _mapped_rows(
    rows: Iterable[Iterable[Any]], *, source_sheet: str
) -> Iterator[tuple[dict[str, Any], str]]:
    header: dict[int, str] | None = None
    for row in rows:
        values = list(row)
        candidate = _header_map(values)
        if "title" in candidate.values() and "view_count" in candidate.values():
            header = candidate
            continue
        if header is None or not any(value not in (None, "") for value in values):
            continue
        mapped = {
            field: values[index] if index < len(values) else None
            for index, field in header.items()
        }
        if _text(mapped.get("title")):
            yield mapped, source_sheet
    if header is None:
        raise _MissingWorkHeader(
            "data file has no recognizable title and view-count header"
        )


def _append_mapped_rows(
    destination: list[tuple[dict[str, Any], str]],
    rows: Iterable[tuple[dict[str, Any], str]],
) -> None:
    for row in rows:
        if len(destination) >= MAX_CREATOR_EXPORT_ROWS:
            raise WorkDataError(
                "creator export exceeds "
                f"the {MAX_CREATOR_EXPORT_ROWS:,}-work-row limit"
            )
        destination.append(row)


def _validate_xlsx_archive(path: Path) -> None:
    try:
        with zipfile.ZipFile(path) as archive:
            members = archive.infolist()
    except (OSError, zipfile.BadZipFile) as exc:
        raise WorkDataError("cannot open XLSX creator export") from exc
    if len(members) > MAX_XLSX_ARCHIVE_MEMBERS:
        raise WorkDataError(
            "XLSX archive exceeds "
            f"the {MAX_XLSX_ARCHIVE_MEMBERS:,}-member limit"
        )
    uncompressed_bytes = sum(member.file_size for member in members)
    if uncompressed_bytes > MAX_XLSX_UNCOMPRESSED_BYTES:
        limit_mib = MAX_XLSX_UNCOMPRESSED_BYTES / (1024 * 1024)
        raise WorkDataError(
            "XLSX archive uncompressed content exceeds "
            f"the {limit_mib:g}-MiB limit"
        )


def read_creator_export(
    path: Path, *, platform: str = "douyin", observed_at: str | None = None
) -> list[dict[str, Any]]:
    suffix = path.suffix.lower()
    mapped: list[tuple[dict[str, Any], str]] = []
    if suffix == ".csv":
        text: str | None = None
        for encoding in ("utf-8-sig", "gb18030"):
            try:
                text = path.read_text(encoding=encoding)
                break
            except UnicodeDecodeError:
                continue
        if text is None:
            raise WorkDataError("CSV must be UTF-8 or GB18030 encoded")
        _append_mapped_rows(
            mapped,
            _mapped_rows(csv.reader(text.splitlines()), source_sheet="csv"),
        )
    elif suffix == ".xlsx":
        try:
            from openpyxl import load_workbook
        except ImportError as exc:  # pragma: no cover - dependency installation failure
            raise WorkDataError("openpyxl is required to import XLSX files") from exc
        _validate_xlsx_archive(path)
        try:
            workbook = load_workbook(path, read_only=True, data_only=True)
        except Exception as exc:
            raise WorkDataError("cannot open XLSX creator export") from exc
        try:
            for sheet in workbook.worksheets:
                try:
                    _append_mapped_rows(
                        mapped,
                        _mapped_rows(
                            sheet.iter_rows(values_only=True),
                            source_sheet=sheet.title,
                        )
                    )
                except _MissingWorkHeader:
                    # Creator-center exports commonly include instruction or
                    # summary sheets before the actual work-data sheet.
                    continue
        finally:
            workbook.close()
    else:
        raise WorkDataError("creator export must use the .xlsx or .csv extension")
    observed = observed_at or utc_now()
    records = [
        normalize_work_snapshot(
            row,
            platform=platform,
            observed_at=observed,
            source_file=path.name,
            source_sheet=sheet,
        )
        for row, sheet in mapped
    ]
    if not records:
        raise WorkDataError("creator export contains no recognizable work rows")
    return records


def load_work_snapshots(path: Path) -> list[dict[str, Any]]:
    if not path.exists():
        return []
    result: list[dict[str, Any]] = []
    with path.open("r", encoding="utf-8-sig") as handle:
        for line_number, line in enumerate(handle, start=1):
            if not line.strip():
                continue
            try:
                value = json.loads(line)
            except json.JSONDecodeError as exc:
                raise WorkDataError(f"{path}:{line_number}: invalid JSON") from exc
            if not isinstance(value, dict):
                raise WorkDataError(f"{path}:{line_number}: record must be an object")
            result.append(normalize_work_snapshot(value))
    return result


def merge_work_snapshots(
    existing: Iterable[Mapping[str, Any]], incoming: Iterable[Mapping[str, Any]]
) -> tuple[list[dict[str, Any]], dict[str, int]]:
    by_identity: dict[tuple[str, str, str], dict[str, Any]] = {}
    existing_count = 0
    duplicate_count = 0
    for source, is_existing in ((existing, True), (incoming, False)):
        for raw in source:
            record = normalize_work_snapshot(raw)
            identity = (
                record["platform"],
                record["work_id"],
                record["observed_at"],
            )
            current = by_identity.get(identity)
            if current is not None:
                if current != record:
                    raise WorkDataError(
                        "conflicting work snapshot for "
                        f"{record['work_id']} at {record['observed_at']}"
                    )
                duplicate_count += 1
                continue
            by_identity[identity] = record
            if is_existing:
                existing_count += 1
    merged = [by_identity[key] for key in sorted(by_identity, key=lambda item: (item[2], item[0], item[1]))]
    return merged, {
        "existing_snapshots": existing_count,
        "incoming_snapshots": len(by_identity) - existing_count,
        "stored_snapshots": len(merged),
        "duplicate_snapshots": duplicate_count,
        "work_count": len({(row["platform"], row["work_id"]) for row in merged}),
    }


def write_work_snapshots(path: Path, records: Iterable[Mapping[str, Any]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    descriptor, temporary_name = tempfile.mkstemp(
        prefix=f".{path.name}.", suffix=".tmp", dir=path.parent
    )
    try:
        with os.fdopen(descriptor, "w", encoding="utf-8", newline="\n") as handle:
            for record in records:
                handle.write(json.dumps(dict(record), ensure_ascii=False, separators=(",", ":")))
                handle.write("\n")
            handle.flush()
            os.fsync(handle.fileno())
        os.replace(temporary_name, path)
    except Exception:
        try:
            os.unlink(temporary_name)
        except FileNotFoundError:
            pass
        raise
