from __future__ import annotations

import csv
from pathlib import Path

import pytest
from openpyxl import Workbook

import admin_app.work_data as work_data
from admin_app.work_data import (
    WorkDataError,
    merge_work_snapshots,
    normalize_work_snapshot,
    read_creator_export,
)


OBSERVED_AT = "2026-09-06T12:00:00+08:00"


def test_csv_creator_export_normalizes_counts_percentages_and_stable_id(
    tmp_path: Path,
) -> None:
    path = tmp_path / "works.csv"
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(
            [
                "作品标题",
                "发布时间",
                "播放量",
                "完播率",
                "5秒完播率",
                "封面点击率",
                "2秒跳出率",
                "平均播放时长",
                "点赞量",
                "作品标签",
            ]
        )
        writer.writerow(
            [
                "低速跟车怎么开",
                "2026-09-01 08:30:00+08:00",
                "1,234",
                "37.5%",
                "68",
                "0.12",
                "22%",
                "18.6",
                "91",
                "#双离合, 驾驶技巧",
            ]
        )

    records = read_creator_export(path, observed_at=OBSERVED_AT)

    assert len(records) == 1
    record = records[0]
    assert record["view_count"] == 1_234
    assert record["like_count"] == 91
    assert record["completion_rate"] == 0.375
    assert record["five_second_completion_rate"] == 0.68
    assert record["cover_click_rate"] == 0.12
    assert record["two_second_bounce_rate"] == 0.22
    assert record["average_watch_seconds"] == 18.6
    assert record["tags"] == ["双离合", "驾驶技巧"]
    assert record["work_id_kind"] == "synthetic"
    assert record["work_id"] == normalize_work_snapshot(
        {
            "title": "低速跟车怎么开",
            "published_at": "2026-09-01T00:30:00Z",
            "observed_at": "2026-09-07T00:00:00Z",
        }
    )["work_id"]


@pytest.mark.parametrize(
    "id_header",
    ["work_id", "作品ID", "视频ID", "item_id", "aweme_id"],
)
def test_creator_export_prefers_platform_work_id(
    tmp_path: Path,
    id_header: str,
) -> None:
    path = tmp_path / "works-with-id.csv"
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow([id_header, "作品标题", "播放量"])
        writer.writerow(["7390123456789012345", "平台 ID 测试", 42])

    record = read_creator_export(path, observed_at=OBSERVED_AT)[0]

    assert record["work_id"] == "7390123456789012345"
    assert record["work_id_kind"] == "platform"


def test_xlsx_creator_export_matches_csv_semantics_and_stable_id(
    tmp_path: Path,
) -> None:
    path = tmp_path / "works.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "作品数据"
    sheet.append(["导出说明", "以下为作品数据"])
    sheet.append(["标题", "发布时间", "播放", "完播率", "评论", "涨粉量"])
    sheet.append(
        [
            "低速跟车怎么开",
            "2026/09/01 08:30:00+08:00",
            1234,
            0.375,
            18,
            7,
        ]
    )
    workbook.save(path)
    workbook.close()

    record = read_creator_export(path, observed_at=OBSERVED_AT)[0]

    assert record["source_file"] == "works.xlsx"
    assert record["source_sheet"] == "作品数据"
    assert record["completion_rate"] == 0.375
    assert record["comment_count"] == 18
    assert record["follower_gain"] == 7
    assert record["work_id"] == normalize_work_snapshot(
        {
            "title": "低速跟车怎么开",
            "published_at": "2026-09-01T08:30:00+08:00",
            "observed_at": OBSERVED_AT,
        }
    )["work_id"]


def test_xlsx_creator_export_skips_auxiliary_sheets_without_known_headers(
    tmp_path: Path,
) -> None:
    path = tmp_path / "works-with-instructions.xlsx"
    workbook = Workbook()
    instructions = workbook.active
    instructions.title = "导出说明"
    instructions.append(["数据口径", "统计结果可能存在延迟"])
    data = workbook.create_sheet("作品数据")
    data.append(["标题", "播放量", "点赞量"])
    data.append(["雨天低速驾驶", 560, 24])
    workbook.save(path)
    workbook.close()

    records = read_creator_export(path, observed_at=OBSERVED_AT)

    assert len(records) == 1
    assert records[0]["title"] == "雨天低速驾驶"
    assert records[0]["view_count"] == 560
    assert records[0]["source_sheet"] == "作品数据"


def test_xlsx_creator_export_rejects_workbook_without_recognizable_data(
    tmp_path: Path,
) -> None:
    path = tmp_path / "instructions-only.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "导出说明"
    sheet.append(["数据口径", "统计结果可能存在延迟"])
    workbook.create_sheet("汇总").append(["合计", 100])
    workbook.save(path)
    workbook.close()

    with pytest.raises(
        WorkDataError,
        match="creator export contains no recognizable work rows",
    ):
        read_creator_export(path, observed_at=OBSERVED_AT)


def test_csv_creator_export_still_requires_recognizable_headers(
    tmp_path: Path,
) -> None:
    path = tmp_path / "instructions.csv"
    path.write_text("数据口径,统计结果可能存在延迟\n", encoding="utf-8")

    with pytest.raises(
        WorkDataError,
        match="data file has no recognizable title and view-count header",
    ):
        read_creator_export(path, observed_at=OBSERVED_AT)


@pytest.mark.parametrize(
    ("limit_name", "message"),
    [
        ("MAX_XLSX_ARCHIVE_MEMBERS", "XLSX archive exceeds"),
        ("MAX_XLSX_UNCOMPRESSED_BYTES", "uncompressed content exceeds"),
    ],
)
def test_xlsx_creator_export_rejects_oversized_archive_before_parsing(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    limit_name: str,
    message: str,
) -> None:
    path = tmp_path / "bounded.xlsx"
    workbook = Workbook()
    workbook.active.append(["标题", "播放量"])
    workbook.active.append(["边界测试", 1])
    workbook.save(path)
    workbook.close()
    monkeypatch.setattr(work_data, limit_name, 1)
    monkeypatch.setattr(
        "openpyxl.load_workbook",
        lambda *args, **kwargs: pytest.fail("openpyxl parsed an unsafe archive"),
    )

    with pytest.raises(WorkDataError, match=message):
        read_creator_export(path, observed_at=OBSERVED_AT)


def test_creator_export_rejects_more_than_maximum_work_rows(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    path = tmp_path / "too-many-works.csv"
    path.write_text(
        "作品标题,播放量\n第一条,1\n第二条,2\n",
        encoding="utf-8",
    )
    monkeypatch.setattr(work_data, "MAX_CREATOR_EXPORT_ROWS", 1)

    with pytest.raises(
        WorkDataError,
        match="creator export exceeds the 1-work-row limit",
    ):
        read_creator_export(path, observed_at=OBSERVED_AT)


def test_xlsx_creator_export_enforces_maximum_work_rows(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    path = tmp_path / "too-many-works.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["作品标题", "播放量"])
    sheet.append(["第一条", 1])
    sheet.append(["第二条", 2])
    workbook.save(path)
    workbook.close()
    monkeypatch.setattr(work_data, "MAX_CREATOR_EXPORT_ROWS", 1)

    with pytest.raises(
        WorkDataError,
        match="creator export exceeds the 1-work-row limit",
    ):
        read_creator_export(path, observed_at=OBSERVED_AT)


def test_work_snapshot_history_deduplicates_exact_observations_only() -> None:
    first = normalize_work_snapshot(
        {
            "work_id": "work-1",
            "title": "第一条作品",
            "observed_at": "2026-09-06T01:00:00Z",
            "view_count": 100,
        }
    )
    later = normalize_work_snapshot(
        {
            **first,
            "observed_at": "2026-09-07T01:00:00Z",
            "view_count": 180,
        }
    )

    merged, stats = merge_work_snapshots([first], [first, later])

    assert [record["observed_at"] for record in merged] == [
        "2026-09-06T01:00:00Z",
        "2026-09-07T01:00:00Z",
    ]
    assert stats == {
        "existing_snapshots": 1,
        "incoming_snapshots": 1,
        "stored_snapshots": 2,
        "duplicate_snapshots": 1,
        "work_count": 1,
    }

    conflicting = {**first, "view_count": 101}
    with pytest.raises(WorkDataError, match="conflicting work snapshot"):
        merge_work_snapshots([first], [conflicting])
