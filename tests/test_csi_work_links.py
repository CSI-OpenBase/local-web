from io import BytesIO
from pathlib import Path
from typing import Any

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from admin_app.csi_work_links import (
    ARCHIVE_HEADERS,
    build_csi_work_archive,
    build_csi_work_archive_rows,
)
from admin_app.local_app import create_local_app
from admin_app.local_config import LocalSettings
from admin_app.local_store import LocalStore


class FakeRunner:
    def start(self) -> None:
        return None

    def submit(self, kind: str, **payload: Any) -> dict[str, Any]:
        raise AssertionError("exporting a work archive must not create a job")


def test_archive_exports_only_video_id_title_and_beijing_publish_time() -> None:
    videos = [
        {
            "video_id": "7390123456789012345",
            "platform": "douyin",
            "title": "离合器保养",
            "last_seen_at": "2026-09-28T08:00:00Z",
            "record": {
                "title": "离合器保养",
                "desc": "离合器保养\n三个检查步骤 #汽车知识",
                "published_at": "2026-09-01T00:30:00Z",
                "author": {"uid": "must-not-export"},
                "cookie": "must-not-export",
            },
        }
    ]

    rows = build_csi_work_archive_rows(videos)
    workbook = load_workbook(BytesIO(build_csi_work_archive(videos)), read_only=True)
    sheet = workbook["视频档案"]

    assert rows == [
        (
            "7390123456789012345",
            "离合器保养\n三个检查步骤 #汽车知识",
            "2026-09-01 08:30:00",
        )
    ]
    assert tuple(cell.value for cell in sheet[1]) == ARCHIVE_HEADERS
    assert tuple(cell.value for cell in sheet[2]) == rows[0]
    assert sheet.max_column == 3


def test_archive_deduplicates_ids_and_ignores_other_platforms() -> None:
    rows = build_csi_work_archive_rows(
        [
            {"video_id": "7390123456789012345", "platform": "douyin"},
            {"video_id": "7390123456789012345", "platform": "douyin"},
            {"video_id": "8490123456789012345", "platform": "tiktok"},
            {"video_id": "", "platform": "douyin"},
        ]
    )

    assert [row[0] for row in rows] == ["7390123456789012345"]


def test_local_settings_downloads_a_csi_work_archive(tmp_path: Path) -> None:
    settings = LocalSettings(
        data_home=tmp_path,
        session_home=tmp_path / ".sessions",
    )
    store = LocalStore(settings.database_path)
    store.upsert_videos(
        [
            {
                "video_id": "7390123456789012350",
                "platform": "douyin",
                "title": "归档作品",
                "video_url": "https://www.douyin.com/video/7390123456789012350",
                "manifest_path": "works/videos/douyin/7390123456789012350/manifest.json",
                "first_seen_at": "2026-09-28T08:00:00Z",
                "last_seen_at": "2026-09-28T09:00:00Z",
                "desc": "归档作品完整描述",
                "published_at": "2026-09-01T00:30:00Z",
            }
        ]
    )

    with TestClient(create_local_app(settings, store=store, runner=FakeRunner())) as client:
        settings_page = client.get("/settings")
        response = client.get("/exports/csi-work-links")

    assert settings_page.status_code == 200
    assert "导出作品档案" in settings_page.text
    assert response.status_code == 200
    assert response.headers["cache-control"] == "no-store"
    assert response.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert 'filename="csi-work-archive-' in response.headers["content-disposition"]

    workbook = load_workbook(BytesIO(response.content), read_only=True)
    sheet = workbook["视频档案"]
    assert tuple(cell.value for cell in sheet[1]) == ARCHIVE_HEADERS
    assert tuple(cell.value for cell in sheet[2]) == (
        "7390123456789012350",
        "归档作品完整描述",
        "2026-09-01 08:30:00",
    )
